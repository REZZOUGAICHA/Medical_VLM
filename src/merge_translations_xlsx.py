"""Merge Hindi/Mandarin translations into the final team deliverable xlsx,
matching the existing annotated_fr/annotated_es column convention.

Recomputes text_to_translate per row directly from the xlsx's own
report/radlex_phrase columns (same priority as prepare_csv_for_translation.py:
report if real, else radlex_phrase, else blank) rather than assuming row
order matches prepped_for_translation.csv -- self-contained, no alignment
assumption.

Hindi note: `--unique-hi` should be unique_texts_hi_v2.txt (the term-
reconstructed version for radlex_phrase rows), not the original
unique_texts_hi.txt -- see reconstruct_radlex_hindi.py. Known limitation:
several of the 35 RadLex terms are still mistranslated by the underlying
model (see docs/process_guide.md) -- this is documented, not silently
shipped as verified-correct.

Usage:
    python src/merge_translations_xlsx.py \
        --xlsx "data/to_translate/final_ALB_dataset(fr+es annotations).xlsx" \
        --unique-hi data/to_translate/unique_texts_hi_v2.txt \
        --unique-zh data/to_translate/unique_texts_zh.txt \
        --unique-en data/to_translate/unique_texts.txt \
        --out "data/to_translate/final_ALB_dataset(fr+es+hi+zh annotations).xlsx"
"""
import argparse
from pathlib import Path

import openpyxl

PLACEHOLDER_VALUES = {"", "not mentioned", "n/a", "na", "none", "nan"}


def pick_text(report, phrase):
    report = (report or "").strip()
    if report and report.lower() not in PLACEHOLDER_VALUES:
        return report
    phrase = (phrase or "").strip()
    if phrase:
        return phrase
    return ""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--unique-en", required=True, type=Path)
    ap.add_argument("--unique-hi", required=True, type=Path)
    ap.add_argument("--unique-zh", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args()

    en_lines = args.unique_en.read_text(encoding="utf-8").splitlines()
    hi_lines = args.unique_hi.read_text(encoding="utf-8").splitlines()
    zh_lines = args.unique_zh.read_text(encoding="utf-8").splitlines()
    assert len(en_lines) == len(hi_lines) == len(zh_lines), (
        f"line count mismatch: en={len(en_lines)} hi={len(hi_lines)} zh={len(zh_lines)}"
    )
    lookup = {en: (hi, zh) for en, hi, zh in zip(en_lines, hi_lines, zh_lines)}
    print(f"Loaded {len(lookup)} unique translated strings")

    wb_in = openpyxl.load_workbook(args.xlsx, read_only=True)
    ws_in = wb_in["Sheet1"]
    rows_iter = ws_in.iter_rows(min_row=1, values_only=True)
    header = list(next(rows_iter))
    idx = {name: i for i, name in enumerate(header)}
    for col in ("report", "radlex_phrase"):
        assert col in idx, f"expected column {col!r} in {args.xlsx}"

    wb_out = openpyxl.Workbook(write_only=True)
    ws_out = wb_out.create_sheet("Sheet1")
    ws_out.append(header + ["annotated_hi", "annotated_zh"])

    n_rows = 0
    n_translated = 0
    n_missing = 0
    for row in rows_iter:
        n_rows += 1
        text = pick_text(row[idx["report"]], row[idx["radlex_phrase"]])
        if not text:
            hi, zh = "", ""
        elif text not in lookup:
            hi, zh = "", ""
            n_missing += 1
        else:
            hi, zh = lookup[text]
            n_translated += 1
        ws_out.append(list(row) + [hi, zh])

    print(f"Rows: {n_rows}, translated: {n_translated} ({n_translated/n_rows*100:.1f}%), missing from lookup: {n_missing}")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
